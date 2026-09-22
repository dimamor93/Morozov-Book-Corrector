
from __future__ import annotations

from pathlib import Path
from docx import Document
from docx.shared import RGBColor, Pt

from .diff import diff


def add_diff(paragraph, original: str, corrected: str):
    """
    Render EXACTLY the corrected sentence.

    Only characters introduced/changed by the correction are red.
    Whitespace is never coloured on its own.
    """
    for kind, text in diff(original, corrected):
        run = paragraph.add_run(text)
        if kind == "changed":
            run.font.color.rgb = RGBColor(255, 0, 0)


def create_report(path: Path, corrections, suspicious, total):
    doc = Document()

    title = doc.add_paragraph()
    run = title.add_run("ОРФОГРАФИЧЕСКИЙ ОТЧЁТ")
    run.bold = True
    run.font.size = Pt(18)

    doc.add_paragraph(f"Всего предложений: {total}")
    doc.add_paragraph(f"Предложений с исправлениями: {len(corrections)}")
    doc.add_paragraph(f"Подозрительных ответов модели: {len(suspicious)}")
    doc.add_paragraph("Исправленные фрагменты выделены красным.")

    for item in sorted(corrections, key=lambda x: x["id"]):
        h = doc.add_paragraph()
        r = h.add_run(f"Предложение № {item['id']}")
        r.bold = True

        p = doc.add_paragraph()
        add_diff(p, item["original"], item["corrected"])

    if suspicious:
        doc.add_page_break()
        h = doc.add_paragraph()
        r = h.add_run("ПОДОЗРИТЕЛЬНЫЕ ОТВЕТЫ МОДЕЛИ")
        r.bold = True

        for item in suspicious:
            doc.add_paragraph(str(item))

    doc.save(path)
