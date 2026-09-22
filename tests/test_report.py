from pathlib import Path
from docx import Document
from docx.shared import RGBColor

from corrector.report import add_diff, create_report


def _read_runs_text_and_colors(docx_path: Path):
    doc = Document(docx_path)
    out = []
    for p in doc.paragraphs:
        for r in p.runs:
            color = None
            try:
                color = r.font.color.rgb
            except Exception:
                color = None
            out.append((r.text, color))
    return out


def test_add_diff_marks_changed_red():
    doc = Document()
    p = doc.add_paragraph()
    add_diff(p, "Он пошол к двери.", "Он пошёл к двери.")
    reds = [t for r in p.runs for t in [r.text]
            if r.font.color.rgb == RGBColor(255, 0, 0)]
    full = "".join(r.text for r in p.runs)
    assert full == "Он пошёл к двери."
    assert any("ё" in t or "о" in t or "ё" in t for t in reds)


def test_add_diff_identical_no_red():
    doc = Document()
    p = doc.add_paragraph()
    add_diff(p, "Всё хорошо.", "Всё хорошо.")
    assert all(r.font.color.rgb is None or r.font.color.rgb != RGBColor(255, 0, 0) for r in p.runs)


def test_create_report_basic(tmp_path: Path):
    out = tmp_path / "report.docx"
    corrections = [
        {"id": 2, "original": "Он пошол.", "corrected": "Он пошёл."},
        {"id": 1, "original": "Привет мир", "corrected": "Привет, мир"},
    ]
    create_report(out, corrections, [], total=10)
    assert out.exists()
    doc = Document(out)
    texts = "\n".join(p.text for p in doc.paragraphs)
    assert "ОРФОГРАФИЧЕСКИЙ ОТЧЁТ" in texts
    assert "Всего предложений: 10" in texts
    assert "Предложение № 1" in texts
    assert "Предложение № 2" in texts
    # сортировка по id: №1 раньше №2
    assert texts.index("Предложение № 1") < texts.index("Предложение № 2")


def test_create_report_with_suspicious(tmp_path: Path):
    out = tmp_path / "rep2.docx"
    create_report(out, [], [{"id": 5, "reason": "large_edit"}], total=3)
    doc = Document(out)
    texts = "\n".join(p.text for p in doc.paragraphs)
    assert "ПОДОЗРИТЕЛЬНЫЕ" in texts


def test_create_report_empty(tmp_path: Path):
    out = tmp_path / "empty.docx"
    create_report(out, [], [], total=0)
    assert out.exists()
    doc = Document(out)
    assert any("Всего предложений: 0" in p.text for p in doc.paragraphs)
