import io
import json
from pathlib import Path
from unittest.mock import MagicMock, patch

import pytest
from fastapi.testclient import TestClient
from docx import Document

import webapp.app as app_mod
from corrector.models import Sentence


@pytest.fixture()
def client(tmp_path, monkeypatch):
    # изолируем jobs и uploads
    app_mod.jobs.clear()
    monkeypatch.setattr(app_mod, "UPLOADS", tmp_path / "uploads")
    (tmp_path / "uploads").mkdir(parents=True, exist_ok=True)
    with TestClient(app_mod.app) as c:
        yield c
    app_mod.jobs.clear()


def _docx_bytes(paragraphs=("Первое. Второе.",)):
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


class _SyncThread:
    def __init__(self, target=None, daemon=None):
        self._target = target

    def start(self):
        self._target()


def _patch_worker_env(monkeypatch, report_content=True):
    monkeypatch.setattr(app_mod, "list_models", lambda: ["qwen:test", "m1"])
    monkeypatch.setattr(
        app_mod, "extract_sentences",
        lambda path: [Sentence(id=1, text="Привет.", paragraph_id=0, paragraph_order=0)],
    )

    class FakePipeline:
        def __init__(self, input_path=None, sentences=None, work_dir=None, config=None, **kw):
            self.input_path = Path(input_path)
            self.work_dir = Path(work_dir)
            self.config = config

        def run(self):
            # имитируем работу pipeline: пишем run_state и отчёт
            self.work_dir.mkdir(parents=True, exist_ok=True)
            (self.work_dir / "run_state.json").write_text(
                json.dumps({"chunks_done": 1, "chunks_total": 1,
                            "eta_seconds": 0, "avg_seconds": 0.1}),
                encoding="utf-8",
            )
            report = self.input_path.parent / self.config["report"]["filename"]
            d = Document()
            d.add_paragraph("fake report")
            d.save(report)

    monkeypatch.setattr(app_mod, "CorrectionPipeline", FakePipeline)
    # синхронный воркер чтобы не было гонок
    monkeypatch.setattr(app_mod.threading, "Thread", _SyncThread)


def test_health(client):
    r = client.get("/api/health")
    assert r.status_code == 200
    assert r.json() == {"status": "ok"}


def test_index_returns_html(client):
    r = client.get("/")
    assert r.status_code == 200
    assert "Book Corrector" in r.text


def test_models_ok(client, monkeypatch):
    monkeypatch.setattr(app_mod, "list_models", lambda: ["a", "b"])
    r = client.get("/api/models")
    assert r.status_code == 200
    assert r.json() == {"models": ["a", "b"]}


def test_models_unavailable(client, monkeypatch):
    def boom():
        raise RuntimeError("no ollama")
    monkeypatch.setattr(app_mod, "list_models", boom)
    r = client.get("/api/models")
    assert r.status_code == 502


def test_create_job_rejects_non_docx(client, monkeypatch):
    monkeypatch.setattr(app_mod, "list_models", lambda: ["qwen:test"])
    r = client.post("/api/jobs", files={"file": ("x.txt", b"hello")}, data={"model": "qwen:test"})
    assert r.status_code == 400


def test_create_job_rejects_unknown_model(client, monkeypatch):
    monkeypatch.setattr(app_mod, "list_models", lambda: ["qwen:test"])
    content = _docx_bytes()
    r = client.post("/api/jobs", files={"file": ("b.docx", content)}, data={"model": "nope"})
    assert r.status_code == 400


def test_create_job_success_and_download(client, monkeypatch):
    _patch_worker_env(monkeypatch)
    content = _docx_bytes()
    r = client.post("/api/jobs", files={"file": ("book.docx", content)}, data={"model": "qwen:test"})
    assert r.status_code == 200
    job_id = r.json()["job_id"]

    st = client.get(f"/api/jobs/{job_id}").json()
    assert st["status"] == "completed"
    assert st["download_available"] is True
    assert st["progress"] == 1
    assert st["total"] == 1
    # чувствительные поля скрыты
    assert "input" not in st and "report" not in st and "cancel_event" not in st
    assert "can_cancel" in st

    dl = client.get(f"/api/jobs/{job_id}/download")
    assert dl.status_code == 200
    assert len(dl.content) > 0


def test_job_status_404(client):
    assert client.get("/api/jobs/nope").status_code == 404


def test_download_not_ready_404(client, monkeypatch):
    monkeypatch.setattr(app_mod, "list_models", lambda: ["qwen:test"])
    # создаём job вручную в статусе running без отчёта
    import threading as th
    app_mod.jobs["j1"] = {"status": "running", "progress": 0, "total": 1,
                          "message": "x", "input": " Kek ", "report": None,
                          "error": None, "cancel_event": th.Event()}
    assert client.get("/api/jobs/j1/download").status_code == 404


def test_download_missing_file_404(client):
    import threading as th
    app_mod.jobs["j2"] = {"status": "completed", "progress": 1, "total": 1,
                          "message": "Готово", "input": "x",
                          "report": "/nonexistent/report.docx",
                          "error": None, "cancel_event": th.Event()}
    assert client.get("/api/jobs/j2/download").status_code == 404


def test_cancel_flow(client):
    import threading as th
    app_mod.jobs["jc"] = {"status": "running", "progress": 0, "total": 2,
                          "message": "run", "input": "x", "report": None,
                          "error": None, "cancel_event": th.Event()}
    r = client.post("/api/jobs/jc/cancel")
    assert r.status_code == 200
    assert r.json()["status"] == "cancelling"
    assert app_mod.jobs["jc"]["cancel_event"].is_set()


def test_cancel_completed_noop(client):
    import threading as th
    app_mod.jobs["jd"] = {"status": "completed", "progress": 1, "total": 1,
                          "message": "Готово", "input": "x", "report": "y",
                          "error": None, "cancel_event": th.Event()}
    r = client.post("/api/jobs/jd/cancel")
    assert r.json()["status"] == "completed"


def test_cancel_404(client):
    assert client.post("/api/jobs/missing/cancel").status_code == 404


def test_worker_error_path(client, monkeypatch):
    monkeypatch.setattr(app_mod, "list_models", lambda: ["qwen:test"])
    monkeypatch.setattr(app_mod.threading, "Thread", _SyncThread)

    def bad_extract(path):
        raise RuntimeError("broken docx")
    monkeypatch.setattr(app_mod, "extract_sentences", bad_extract)

    content = _docx_bytes()
    r = client.post("/api/jobs", files={"file": ("e.docx", content)}, data={"model": "qwen:test"})
    job_id = r.json()["job_id"]
    st = client.get(f"/api/jobs/{job_id}").json()
    assert st["status"] == "error"
